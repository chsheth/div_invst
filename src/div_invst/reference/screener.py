"""Fidelity stock screener workbook (screener_results.xlsx, built by scripts/merge_screener.py).

Five tabs joined on Symbol. Column meanings, units and fill rates: docs/research/fidelity_screener.md.
Conversions: percent columns become fractions (6.8 -> 0.068) like the DRIP loader; dollar columns stay in $ millions
(`_m` suffix) and market cap becomes dollars; "--" and blanks become None; categorical ratings stay strings.
Standard deviation and the TTM cash flow growth are already fractions; ratios and multiples are unchanged.
"""
import re
from pathlib import Path

from openpyxl import load_workbook

# kinds: str, cat (string, "--" is unknown), num, pct (percent -> fraction), mcap ("$2.83B" -> dollars), bool (Yes/No)
COLUMNS = {
    "Search Criteria": {
        "Company Name": ("company", "str"), "Security Type": ("security_type", "str"), "Security Price": ("price", "num"),
        "Dividend Yield": ("div_yield", "pct"), "Dividend": ("dividend", "num"),
        "Dividend Payout % (Last Quarter)": ("payout", "pct"), "Dividend Growth Rate (5 Year Avg)": ("dgr_5y", "pct"),
        "Dividend Growth Rate (IAD to Prior Yr)": ("dgr_iad_1y", "pct"),
        "Dividend Coverage (EPS TTM/IAD)": ("div_coverage_ttm", "num"), "Dividend Coverage (EPS Next Yr/IAD)": ("div_coverage_next", "num"),
        "Earnings Yield (TTM)": ("earnings_yield", "pct"), "Total Cash": ("cash_m", "num"), "Revenue (TTM)": ("revenue_m", "num"),
        "Jefferson Research": ("jefferson", "cat"), "Price/Book Ratio": ("pb", "num"), "Price/Sales Ratio": ("ps", "num"),
        "Price/Cash Flow Ratio": ("pcf", "num"), "Total Debt/Equity (TTM)": ("debt_to_equity", "pct"),
        "Free Cash Flow": ("fcf_m", "num"), "Return on Equity (TTM)": ("roe", "pct"), "Return on Assets (TTM)": ("roa", "pct"),
        "EBITD Margin (TTM)": ("ebitda_margin", "pct"), "Gross Profit Margin (TTM)": ("gross_margin", "pct"),
        "Forward EPS Long Term Growth (3-5 Yrs)": ("eps_growth_lt_est", "pct"), "Book Value Growth (5 Yr average)": ("book_value_growth_5y", "pct"),
        "P/E (Next Year's Estimate)": ("pe_next", "num"), "Current Ratio (Most Recent Qtr)": ("current_ratio", "num"),
        "Quick Ratio (Most Recent Qtr)": ("quick_ratio", "num"), "Debt to Capital Ratio (Most Recent Qtr)": ("debt_to_capital", "pct"),
        "Interest Coverage (Most Recent Qtr)": ("interest_coverage", "num"), "% Revenue to RD (Last FY)": ("rd_pct_revenue", "pct"),
        "Net Insider Shares Bought": ("insider_buy_shares_m", "num"),
        "Revenue Growth (5 Yrs)": ("revenue_growth_5y", "pct"), "Revenue Growth (3 Yrs)": ("revenue_growth_3y", "pct"),
        "Revenue Growth (TTM vs. Prior TTM)": ("revenue_growth_ttm", "pct"),
        "Cash Flow Growth Rate (3 Yr)": ("cf_growth_3y", "pct"), "Cash Flow Growth Rate (5 Yr)": ("cf_growth_5y", "pct"),
        "Cash Flow Growth Rate (TTM vs. Prior TTM)": ("cf_growth_ttm", "num"),
        "EPS Growth (3 Year History)": ("eps_growth_3y", "pct"), "EPS Growth (5 Year Historical)": ("eps_growth_5y", "pct"),
        "EPS Growth (Proj Next Yr vs. This Yr)": ("eps_growth_next", "pct"),
        "EPS Growth (Proj This Yr vs. Last Yr)": ("eps_growth_this", "pct"),
        "Beta (3 Yr Annualized)": ("beta_3y", "num"), "Beta (5 Yr Annualized)": ("beta_5y", "num"), "Beta (10 Yr Annualized)": ("beta_10y", "num"),
        "Total Return (3 Yr Annualized)": ("total_return_3y", "pct"), "Total Return (5 Yr Annualized)": ("total_return_5y", "pct"),
        "Total Return (10 Yr Annualized)": ("total_return_10y", "pct"),
        "Standard Deviation (3 Yr Annualized)": ("std_3y", "num"), "Standard Deviation (5 Yr Annualized)": ("std_5y", "num"),
        "Standard Deviation (10 Yr Annualized)": ("std_10y", "num"),
        "Zacks Investment Research": ("zacks", "cat"), "Trading Central": ("trading_central", "cat"),
        "Argus Analyst Opinions": ("argus_rating", "cat"), "McLean Capital Management": ("mclean", "cat"),
        "Shares Outstanding": ("shares_out_m", "num"), "Profit Margin (TTM)": ("profit_margin", "pct"),
        "Return on Investment (TTM)": ("roi", "pct"), "Operating Margin (5 Year Avg)": ("op_margin_5y", "pct"),
        "Operating Margin (TTM/5 Year Avg.)": ("op_margin_trend", "num"), "SG&A Expense / Net Sales": ("sga_to_sales", "pct"),
    },
    "Basic Facts": {
        "Volume (90 Day Avg)": ("volume_90d_m", "num"), "Market Capitalization": ("market_cap", "mcap"),
        "Company Headquarters Location": ("country", "str"), "Sector": ("sector", "cat"), "Industry": ("industry", "cat"),
        "Optionable": ("optionable", "bool"),
    },
    "Performance & Volatility": {
        "Price Performance (52 Weeks)": ("perf_52w", "pct"), "Total Return (1 Yr Annualized)": ("total_return_1y", "pct"),
        "Beta (1 Year Annualized)": ("beta_1y", "num"), "Standard Deviation (1 Yr Annualized)": ("std_1y", "num"),
    },
    "Valuation, Growth & Ownership": {
        "S&P Global Market Intelligence Valuation": ("sp_valuation", "num"), "S&P Global Market Intelligence Quality": ("sp_quality", "num"),
        "S&P Global Market Intelligence Growth Stability": ("sp_growth_stability", "num"),
        "S&P Global Market Intelligence Financial Health": ("sp_financial_health", "num"),
        "P/E (Price/TTM Earnings)": ("pe", "num"), "PEG Ratio": ("peg", "num"),
        "Institutional Ownership": ("inst_ownership", "pct"), "Institutional Ownership (Last vs. Prior Qtr)": ("inst_ownership_change", "pct"),
    },
    "Analyst Opinions": {
        "Equity Summary Score (ESS) from LSEG StarMine": ("ess", "cat"), "Equity Summary Score Change (1 Month)": ("ess_change", "cat"),
        "LSEG I/B/E/S Estimates": ("ibes_mean", "num"), "MSCI Overall Environmental, Social & Governance Rating": ("msci_esg", "cat"),
        "MSCI Environmental Score": ("msci_e", "cat"), "MSCI Social Score": ("msci_s", "cat"), "MSCI Governance Score": ("msci_g", "cat"),
    },
}
# Columns that repeat on another tab and are ignored there: Company Name (kept from Search Criteria), Security Price,
# Dividend Yield and "EPS Growth (Proj This Yr vs. Last Yr)" on the Valuation tab.

_MCAP = re.compile(r"^\$?([\d,.]+)\s*([KMBT]?)$")
_MCAP_MULT = {"": 1, "K": 1e3, "M": 1e6, "B": 1e9, "T": 1e12}
_UNKNOWN = {"", "--", "-", "N/A", "n/a", "NM"}


def clean_value(raw, kind: str):
    """One workbook cell to its typed value (see module docstring); anything unknown or unparseable is None."""
    if raw is None:
        return None
    if isinstance(raw, str):
        s = raw.strip()
        if s in _UNKNOWN:
            return None
    else:
        s = raw
    if kind in ("str", "cat"):
        return str(s)
    if kind == "bool":
        return str(s).lower().startswith("yes")
    if kind == "mcap":
        m = _MCAP.match(str(s).replace(" ", ""))
        return float(m[1].replace(",", "")) * _MCAP_MULT[m[2]] if m else None
    try:
        v = float(str(s).replace(",", "").replace("%", "")) if isinstance(s, str) else float(s)
    except ValueError:
        return None
    return round(v / 100, 8) if kind == "pct" else v


def load_screener(path: Path) -> dict[str, dict]:
    """Screener workbook keyed by ticker. Raises ValueError if a tab is missing or its Symbol column is absent.
    Columns not in COLUMNS are ignored; mapped columns missing from a tab are left out of the rows."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, dict] = {}
    try:
        for tab, mapping in COLUMNS.items():
            if tab not in wb.sheetnames:
                raise ValueError(f"{path.name} has no tab {tab!r}")
            rows = wb[tab].iter_rows(values_only=True)
            header = next(rows)
            if "Symbol" not in header:
                raise ValueError(f"{path.name} tab {tab!r} has no Symbol column")
            idx = {h: i for i, h in enumerate(header)}
            for row in rows:
                sym = row[idx["Symbol"]]
                if not sym:
                    continue
                rec = out.setdefault(str(sym).strip(), {"ticker": str(sym).strip()})
                for col, (name, kind) in mapping.items():
                    if col in idx and rec.get(name) is None:
                        rec[name] = clean_value(row[idx[col]], kind)
    finally:
        wb.close()
    return out
